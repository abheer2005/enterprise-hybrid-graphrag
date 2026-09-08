from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".txt",
}

ENTERPRISE_METADATA_FIELDS = {
    "title", "document_version", "status", "effective_from",
    "effective_to", "supersedes", "owner", "classification", "language",
}


def _sha256(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _sidecar_metadata(file_path: Path) -> dict[str, Any]:
    """Read optional, generic governance metadata from ``file.ext.metadata.json``."""
    sidecar = file_path.with_name(file_path.name + ".metadata.json")
    if not sidecar.exists():
        return {}
    payload = json.loads(sidecar.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError(f"Metadata sidecar must contain an object: {sidecar}")
    return {key: payload[key] for key in ENTERPRISE_METADATA_FIELDS if key in payload}


def build_document_metadata(file_path: Path, corpus_root: Path | None = None) -> dict[str, Any]:
    """Create stable identity, provenance, version, and lifecycle metadata."""
    resolved = file_path.resolve()
    root = corpus_root.resolve() if corpus_root else resolved.parent
    try:
        relative_path = resolved.relative_to(root).as_posix()
    except ValueError:
        relative_path = resolved.name
    stat = resolved.stat()
    content_hash = _sha256(resolved)
    metadata = {
        "source": resolved.name,
        "file_path": str(resolved),
        "relative_path": relative_path,
        "source_uri": relative_path,
        "document_id": hashlib.sha256(relative_path.casefold().encode("utf-8")).hexdigest()[:24],
        "content_hash": content_hash,
        "document_version": content_hash[:12],
        "status": "current",
        "modified_at": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        "file_size_bytes": stat.st_size,
    }
    metadata.update(_sidecar_metadata(resolved))
    return metadata


def extract_pdf(file_path: Path) -> list[dict[str, Any]]:
    """Extract text page-by-page from a PDF."""

    records = []

    with fitz.open(file_path) as document:
        for page_number, page in enumerate(document, start=1):
            text = page.get_text("text").strip()

            records.append(
                {
                    "text": text,
                    "metadata": {
                        "source": file_path.name,
                        "file_path": str(file_path.resolve()),
                        "file_type": "pdf",
                        "page": page_number,
                        "total_pages": len(document),
                    },
                }
            )

    return records


def extract_docx(file_path: Path) -> list[dict[str, Any]]:
    """Extract paragraphs from a DOCX document."""

    document = Document(file_path)

    text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    return [
        {
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_path": str(file_path.resolve()),
                "file_type": "docx",
            },
        }
    ]


def extract_pptx(file_path: Path) -> list[dict[str, Any]]:
    """Extract text slide-by-slide from a PowerPoint."""

    presentation = Presentation(file_path)
    records = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                slide_text.append(shape.text.strip())

        records.append(
            {
                "text": "\n".join(slide_text),
                "metadata": {
                    "source": file_path.name,
                    "file_path": str(file_path.resolve()),
                    "file_type": "pptx",
                    "slide": slide_number,
                    "total_slides": len(presentation.slides),
                },
            }
        )

    return records


def extract_xlsx(file_path: Path) -> list[dict[str, Any]]:
    """Extract cell values sheet-by-sheet from an Excel workbook."""

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    records = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [
                str(value).strip()
                for value in row
                if value is not None and str(value).strip()
            ]

            if values:
                rows.append(" | ".join(values))

        records.append(
            {
                "text": "\n".join(rows),
                "metadata": {
                    "source": file_path.name,
                    "file_path": str(file_path.resolve()),
                    "file_type": "xlsx",
                    "sheet": sheet.title,
                },
            }
        )

    workbook.close()

    return records


def extract_txt(file_path: Path) -> list[dict[str, Any]]:
    """Extract a plain text file."""

    text = file_path.read_text(
        encoding="utf-8",
        errors="ignore",
    )

    return [
        {
            "text": text,
            "metadata": {
                "source": file_path.name,
                "file_path": str(file_path.resolve()),
                "file_type": "txt",
            },
        }
    ]


def load_document(file_path: Path, corpus_root: Path | None = None) -> list[dict[str, Any]]:
    """Route a document to the correct extractor."""

    extension = file_path.suffix.lower()

    if extension == ".pdf":
        records = extract_pdf(file_path)

    if extension == ".docx":
        records = extract_docx(file_path)

    if extension == ".pptx":
        records = extract_pptx(file_path)

    if extension == ".xlsx":
        records = extract_xlsx(file_path)

    if extension == ".txt":
        records = extract_txt(file_path)

    if extension in SUPPORTED_EXTENSIONS:
        enterprise_metadata = build_document_metadata(file_path, corpus_root)
        for record in records:
            record["metadata"].update(enterprise_metadata)
        return records

    raise ValueError(
        f"Unsupported file type: {extension}"
    )


def load_directory(directory: Path) -> list[dict[str, Any]]:
    """Discover and load supported documents recursively."""

    all_records = []

    files = sorted(
        path
        for path in directory.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    print(f"\nFound {len(files)} supported document(s).\n")

    for index, file_path in enumerate(files, start=1):
        print(f"[{index}/{len(files)}] Loading: {file_path.name}")

        try:
            records = load_document(file_path, corpus_root=directory)

            non_empty_records = [
                record
                for record in records
                if record["text"].strip()
            ]

            all_records.extend(non_empty_records)

            print(
                f"    Extracted {len(non_empty_records)} "
                f"non-empty section(s)."
            )

        except Exception as exc:
            print(f"    ERROR: {exc}")

    return all_records
